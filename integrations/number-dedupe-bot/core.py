"""Independent number library and lossless-row filtering for the Telegram bot."""

from __future__ import annotations

import csv
import io
import json
import os
import re
import sqlite3
import tempfile
import uuid
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import phonenumbers
from babel import Locale
from openpyxl import Workbook, load_workbook

MAX_INPUT_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 80
MAX_OUTPUT_BYTES = 45 * 1024 * 1024
PHONE_HEADERS = {
    "号码", "手机号", "手机号码", "电话号码", "电话", "客户号码", "客户手机号",
    "phone", "phonenumber", "mobile", "mobilenumber", "number", "tel", "telephone",
}
COUNTRY_HEADERS = {"国家", "国家地区", "国别", "地区", "country", "countryregion", "region"}
UNKNOWN_COUNTRY = "未识别国家"
TEMPLATE_HEADERS = ["编号", "客户电话号码", "客户姓名", "损失金额", "发生欺诈的平台", "其他客户信息", "地址"]
TEMPLATE_LABELS = {
    "客户电话号码": "客户电话号码", "电话号码": "客户电话号码", "手机号": "客户电话号码",
    "客户姓名": "客户姓名", "姓名": "客户姓名", "损失金额": "损失金额",
    "发生欺诈的平台": "发生欺诈的平台", "欺诈平台": "发生欺诈的平台",
    "其他客户信息": "其他客户信息", "地址": "地址", "国家": "国家",
}
TEMPLATE_START = re.compile(r"^\s*编号\s*[:：]?\s*(\S+)\s*$")
TEMPLATE_FIELD = re.compile(r"^\s*([\u4e00-\u9fffA-Za-z][\u4e00-\u9fffA-Za-z0-9 /（）()_-]{0,35})\s*[:：]\s*(.*)$")


def _country_names() -> dict[str, str]:
    result: dict[str, str] = {}
    for language in ("zh", "en"):
        for code, name in Locale.parse(language).territories.items():
            if len(code) == 2:
                result[name.strip().casefold()] = code
                result[code.casefold()] = code
    result.update({"usa": "US", "uk": "GB", "u.s.a.": "US"})
    return result


COUNTRY_NAMES = _country_names()


class ImportErrorForUser(ValueError):
    pass


@dataclass
class Sheet:
    name: str
    headers: list[object]
    rows: list[list[object]]
    phone_index: int
    infer_bare_country: bool = False


@dataclass
class Result:
    total_rows: int
    invalid_rows: int
    repeated_rows: int
    existing_rows: int
    new_rows: int
    output_files: list[tuple[str, Path]]


def normalize_phone(value: object) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        value = int(value)
    if isinstance(value, (int, float)) and len(str(value)) > 15:
        # Excel cannot reliably represent numeric cells longer than 15 digits.
        return None
    raw = str(value).strip()
    if raw.startswith("'"):
        raw = raw[1:]
    if not re.fullmatch(r"\+?[\d\s().\-–—]+", raw):
        return None
    digits = re.sub(r"\D", "", raw)
    if raw.startswith("00") and len(digits) > 8:
        digits = digits[2:]
    return digits if 5 <= len(digits) <= 20 else None


def _header_key(value: object) -> str:
    return re.sub(r"[\s_\-:：()（）/]", "", str(value or "")).lower()


def _bare_international_region(raw: str) -> str | None:
    digits = normalize_phone(raw)
    if not digits or len(digits) < 10:
        return None
    try:
        number = phonenumbers.parse("+" + digits)
        if phonenumbers.is_valid_number(number):
            region = phonenumbers.region_code_for_number(number)
            return region if region and region != "001" else None
    except phonenumbers.NumberParseException:
        pass
    return None


def classify_country(row: list[object], sheet: Sheet) -> str:
    for index, header in enumerate(sheet.headers):
        if _header_key(header) in COUNTRY_HEADERS and index < len(row):
            value = str(row[index] or "").strip()
            if value and value.casefold() not in ("-", "unknown", "n/a", "未知", "未填写"):
                return COUNTRY_NAMES.get(value.casefold(), value[:80])
    raw = str(row[sheet.phone_index] or "").strip()
    if not (raw.startswith("+") or raw.startswith("00")):
        if sheet.infer_bare_country:
            return _bare_international_region(raw) or UNKNOWN_COUNTRY
        return UNKNOWN_COUNTRY
    digits = normalize_phone(raw)
    if not digits:
        return UNKNOWN_COUNTRY
    try:
        number = phonenumbers.parse("+" + digits)
        if phonenumbers.is_possible_number(number):
            region = phonenumbers.region_code_for_number(number)
            if region and region != "001":
                return region
            regions = phonenumbers.region_codes_for_country_code(number.country_code)
            if len(regions) == 1 and regions[0] != "001":
                return regions[0]
    except phonenumbers.NumberParseException:
        pass
    return UNKNOWN_COUNTRY


def number_key(row: list[object], sheet: Sheet) -> tuple[str, str, str] | None:
    raw = row[sheet.phone_index]
    digits = normalize_phone(raw)
    if not digits:
        return None
    country = classify_country(row, sheet)
    text = str(raw).strip().lstrip("'")
    if text.startswith(("+", "00")):
        return "intl:" + digits, digits, country
    has_country_column = any(_header_key(header) in COUNTRY_HEADERS for header in sheet.headers)
    if sheet.infer_bare_country and not has_country_column and _bare_international_region(text):
        return "intl:" + digits, digits, country
    if len(country) == 2 and country.isalpha():
        try:
            parsed = phonenumbers.parse(text, country)
            if phonenumbers.is_valid_number(parsed):
                international = re.sub(r"\D", "", phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164))
                return "intl:" + international, digits, country
        except phonenumbers.NumberParseException:
            pass
    return "local:" + country + ":" + digits, digits, country


def _sheet_from_rows(name: str, incoming: list[list[object]], column: str | None) -> Sheet:
    rows = [row for row in incoming if any(value is not None and str(value).strip() for value in row)]
    if not rows:
        raise ImportErrorForUser(f"工作表「{name}」没有内容。")
    if len(rows) > MAX_ROWS + 1:
        raise ImportErrorForUser(f"工作表「{name}」超过 {MAX_ROWS} 行，请拆成多个文件。")
    if max(map(len, rows)) > MAX_COLUMNS:
        raise ImportErrorForUser(f"工作表「{name}」列数过多。")
    first = rows[0]
    wanted = _header_key(column) if column else None
    matches = [i for i, value in enumerate(first) if _header_key(value) == wanted] if wanted else [
        i for i, value in enumerate(first) if _header_key(value) in PHONE_HEADERS
    ]
    if len(matches) == 1:
        headers = first
        data_rows = rows[1:]
        index = matches[0]
    elif len(matches) > 1:
        raise ImportErrorForUser(f"工作表「{name}」有多个号码列，请用 /column 列名 指定。")
    elif len(first) == 1 and normalize_phone(first[0]):
        headers = ["号码"]
        data_rows = rows
        index = 0
    else:
        raise ImportErrorForUser(f"工作表「{name}」找不到号码列。请把列名改为「号码」，或发送 /column 列名。")
    width = max(len(headers), max((len(row) for row in data_rows), default=0))
    headers = list(headers) + [f"字段{i + 1}" for i in range(len(headers), width)]
    return Sheet(name, headers, [list(row) + [None] * (width - len(row)) for row in data_rows], index)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            text = data.decode(encoding)
            if "\x00" not in text:
                return text
        except UnicodeDecodeError:
            pass
    raise ImportErrorForUser("文本编码无法识别，请保存为 UTF-8、UTF-16 或 GB18030 后重试。")


def looks_like_record_template(text: str) -> bool:
    return bool(re.search(r"(?m)^\s*(?:客户)?(?:电话(?:号码)?|手机号)\s*[:：]", text))


def _parse_record_template(text: str) -> Sheet:
    records: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    active_field: str | None = None
    extra_headers: list[str] = []

    def finish() -> None:
        nonlocal current
        if current is not None and any(current.values()):
            records.append(current)
        current = None

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line or re.fullmatch(r"\d{1,2}:\d{2}", line):
            continue
        start = TEMPLATE_START.match(line)
        if start:
            finish()
            current = {"编号": start.group(1)}
            active_field = None
            continue
        field = TEMPLATE_FIELD.match(line)
        if field:
            key = TEMPLATE_LABELS.get(_header_key(field.group(1)), field.group(1).strip())
            value = field.group(2).strip()
            if key == "客户电话号码" and current is not None and current.get(key):
                finish()
            if current is None:
                current = {}
            if key not in TEMPLATE_HEADERS and key not in extra_headers:
                extra_headers.append(key)
            if current.get(key) and value:
                current[key] += "\n" + value
            else:
                current[key] = value
            active_field = key
            continue
        if current is not None:
            key = active_field or "其他文本"
            if key not in TEMPLATE_HEADERS and key not in extra_headers:
                extra_headers.append(key)
            current[key] = (current.get(key, "") + "\n" + line).strip()
    finish()
    if not records:
        raise ImportErrorForUser("没有识别到客户资料记录。")
    if len(records) > MAX_ROWS:
        raise ImportErrorForUser(f"客户资料超过 {MAX_ROWS} 条，请拆成多个文件。")
    headers = TEMPLATE_HEADERS + extra_headers
    rows = [[record.get(header, "") for header in headers] for record in records]
    return Sheet("客户资料", headers, rows, headers.index("客户电话号码"), infer_bare_country=True)


def _parse_text(data: bytes, filename: str, column: str | None) -> list[Sheet]:
    decoded = _decode_text(data)
    if looks_like_record_template(decoded):
        return [_parse_record_template(decoded)]
    lines = [line.strip() for line in decoded.splitlines() if line.strip()]
    if not lines:
        raise ImportErrorForUser("文件里没有号码。")
    if filename.lower().endswith(".txt") and all(normalize_phone(line) for line in lines):
        return [Sheet("去重结果", ["号码"], [[line] for line in lines], 0)]
    sample = decoded[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|，；")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if "\t" in sample else ","
    try:
        rows = list(csv.reader(io.StringIO(decoded, newline=""), delimiter=delimiter))
    except csv.Error as exc:
        raise ImportErrorForUser("文本表格格式无法识别，请检查分隔符和引号。") from exc
    return [_sheet_from_rows("去重结果", rows, column)]


def parse_file(data: bytes, filename: str, column: str | None = None) -> list[Sheet]:
    if not data or len(data) > MAX_INPUT_BYTES:
        raise ImportErrorForUser("文件为空或超过 20 MB，请拆分后上传。")
    suffix = Path(filename).suffix.lower()
    if suffix in (".txt", ".csv"):
        return _parse_text(data, filename, column)
    if suffix != ".xlsx":
        raise ImportErrorForUser("目前支持 .xlsx、.csv 和 .txt 文件。")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ImportErrorForUser("Excel 解压后过大，请拆分后上传。")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            sheets = []
            for worksheet in workbook.worksheets:
                if worksheet.max_row == 0:
                    continue
                if worksheet.max_row > MAX_ROWS + 1 or worksheet.max_column > MAX_COLUMNS:
                    raise ImportErrorForUser(f"工作表「{worksheet.title}」过大，请拆分后上传。")
                values = [list(row) for row in worksheet.iter_rows(values_only=True)]
                if any(any(value is not None and str(value).strip() for value in row) for row in values):
                    sheets.append(_sheet_from_rows(worksheet.title, values, column))
            if not sheets:
                raise ImportErrorForUser("Excel 里没有号码。")
            if sum(len(sheet.rows) for sheet in sheets) > MAX_ROWS:
                raise ImportErrorForUser(f"Excel 总行数超过 {MAX_ROWS} 行，请拆成多个文件。")
            return sheets
        finally:
            workbook.close()
    except (zipfile.BadZipFile, OSError, KeyError) as exc:
        raise ImportErrorForUser("Excel 文件损坏或格式不正确。") from exc


def _write_result(sheets: list[Sheet], selected: list[list[list[object]]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet, rows in zip(sheets, selected):
        output = workbook.create_sheet(sheet.name[:31])
        for row in [sheet.headers, *rows]:
            output.append(row)
            # openpyxl's append returns None; set text cells after append so input
            # strings beginning with '=' never become executable spreadsheet formulas.
            current = output.max_row
            for index, value in enumerate(row, start=1):
                cell = output.cell(current, index)
                if isinstance(value, str):
                    cell.data_type = "s"
                if index == sheet.phone_index + 1:
                    cell.number_format = "@"
        output.freeze_panes = "A2"
        output.auto_filter.ref = output.dimensions
    stream = io.BytesIO()
    workbook.save(stream)
    result = stream.getvalue()
    if len(result) > MAX_OUTPUT_BYTES:
        raise ImportErrorForUser("去重结果超过电报发送上限，请拆成多个文件。")
    return result


def open_database(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.execute("PRAGMA busy_timeout=5000")
    connection.execute("CREATE TABLE IF NOT EXISTS known_numbers (number_key TEXT PRIMARY KEY, number TEXT NOT NULL, country TEXT NOT NULL, record_json TEXT NOT NULL, created_at TEXT NOT NULL)")
    connection.execute("CREATE TABLE IF NOT EXISTS jobs (id TEXT PRIMARY KEY, chat_id INTEGER NOT NULL, user_id INTEGER NOT NULL, manifest_json TEXT NOT NULL, created_at TEXT NOT NULL)")
    connection.execute("CREATE TABLE IF NOT EXISTS imports (update_id INTEGER PRIMARY KEY, chat_id INTEGER NOT NULL, user_id INTEGER NOT NULL, mode TEXT NOT NULL, total_rows INTEGER NOT NULL, invalid_rows INTEGER NOT NULL, repeated_rows INTEGER NOT NULL, existing_rows INTEGER NOT NULL, new_rows INTEGER NOT NULL, manifest_json TEXT NOT NULL, delivery_index INTEGER NOT NULL DEFAULT 0)")
    connection.execute("CREATE TABLE IF NOT EXISTS user_settings (chat_id INTEGER NOT NULL, user_id INTEGER NOT NULL, mode TEXT, mode_expires_at INTEGER, column_name TEXT, PRIMARY KEY(chat_id, user_id))")
    connection.commit()
    os.chmod(path, 0o600)
    return connection


def latest_result(connection: sqlite3.Connection, chat_id: int, user_id: int) -> list[tuple[str, Path]]:
    row = connection.execute(
        "SELECT manifest_json FROM jobs WHERE chat_id=? AND user_id=? ORDER BY rowid DESC LIMIT 1",
        (chat_id, user_id),
    ).fetchone()
    return [(country, Path(path)) for country, path in json.loads(row[0])] if row else []


def process_file(
    connection: sqlite3.Connection,
    data: bytes,
    filename: str,
    mode: str,
    output_dir: Path,
    chat_id: int,
    user_id: int,
    column: str | None = None,
    update_id: int | None = None,
) -> Result:
    if mode not in ("base", "check"):
        raise ValueError("Unknown import mode")
    sheets = parse_file(data, filename, column)
    identities = [[number_key(row, sheet) for row in sheet.rows] for sheet in sheets]
    frequencies = Counter(identity[0] for group in identities for identity in group if identity)
    total = sum(len(sheet.rows) for sheet in sheets)
    invalid = repeated = existing = added = 0
    selected_by_country: dict[str, list[list[list[object]]]] = {}
    output_files: list[tuple[str, Path]] = []
    now = datetime.now(timezone.utc).isoformat()
    connection.execute("BEGIN IMMEDIATE")
    try:
        if update_id is not None:
            previous = connection.execute(
                "SELECT total_rows, invalid_rows, repeated_rows, existing_rows, new_rows, manifest_json FROM imports WHERE update_id=?",
                (update_id,),
            ).fetchone()
            if previous:
                connection.commit()
                return Result(*previous[:5], [(country, Path(path)) for country, path in json.loads(previous[5])])
        seen_in_base: set[str] = set()
        for sheet_index, sheet in enumerate(sheets):
            for row, identity in zip(sheet.rows, identities[sheet_index]):
                if not identity:
                    invalid += 1
                    continue
                key, phone, country = identity
                if mode == "check" and frequencies[key] > 1:
                    repeated += 1
                    continue
                if mode == "base":
                    if key in seen_in_base:
                        repeated += 1
                        continue
                    seen_in_base.add(key)
                if connection.execute("SELECT 1 FROM known_numbers WHERE number_key=?", (key,)).fetchone():
                    existing += 1
                    continue
                record = list(zip((str(header) for header in sheet.headers), row))
                connection.execute(
                    "INSERT INTO known_numbers(number_key, number, country, record_json, created_at) VALUES (?, ?, ?, ?, ?)",
                    (key, phone, country, json.dumps(record, ensure_ascii=False, default=str), now),
                )
                added += 1
                if mode == "check":
                    output_row = list(row)
                    output_row[sheet.phone_index] = str(row[sheet.phone_index])
                    selected_by_country.setdefault(country, [[] for _ in sheets])[sheet_index].append(output_row)
        if mode == "check":
            output_dir.mkdir(mode=0o700, parents=True, exist_ok=True)
            os.chmod(output_dir, 0o700)
            job_id = uuid.uuid4().hex
            if not selected_by_country:
                selected_by_country["无新增"] = [[] for _ in sheets]
            for index, (country, selected) in enumerate(sorted(selected_by_country.items()), start=1):
                output_bytes = _write_result(sheets, selected)
                output_path = output_dir / f"{job_id}-{index}.xlsx"
                fd, temporary = tempfile.mkstemp(prefix=".partial-", dir=output_dir)
                try:
                    with os.fdopen(fd, "wb") as file:
                        file.write(output_bytes)
                        file.flush()
                        os.fsync(file.fileno())
                    os.replace(temporary, output_path)
                finally:
                    if os.path.exists(temporary):
                        os.unlink(temporary)
                output_files.append((country, output_path))
            connection.execute(
                "INSERT INTO jobs(id, chat_id, user_id, manifest_json, created_at) VALUES (?, ?, ?, ?, ?)",
                (job_id, chat_id, user_id, json.dumps([(country, str(path)) for country, path in output_files]), now),
            )
        if update_id is not None:
            connection.execute(
                "INSERT INTO imports(update_id, chat_id, user_id, mode, total_rows, invalid_rows, repeated_rows, existing_rows, new_rows, manifest_json) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (update_id, chat_id, user_id, mode, total, invalid, repeated, existing, added, json.dumps([(country, str(path)) for country, path in output_files])),
            )
        connection.commit()
    except Exception:
        connection.rollback()
        for _, output_path in output_files:
            output_path.unlink(missing_ok=True)
        raise
    return Result(total, invalid, repeated, existing, added, output_files)
