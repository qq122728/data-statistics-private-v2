import io
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "integrations" / "number-dedupe-bot"))
from core import ImportErrorForUser, latest_result, normalize_phone, open_database, process_file
from bot import Bot


def xlsx(headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def rows_from(path):
    workbook = load_workbook(path, read_only=True)
    return list(workbook.active.values)


class NumberDedupeBotTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.db = open_database(self.root / "numbers.sqlite3")
        self.addCleanup(self.db.close)

    def process(self, data, filename, mode, update_id):
        return process_file(self.db, data, filename, mode, self.root / "results", -100, 7, update_id=update_id)

    def test_base_then_check_excludes_all_file_duplicates_and_existing(self):
        base = self.process(b"+14155552671\n+442079460958\n", "base.txt", "base", 1)
        self.assertEqual(base.new_rows, 2)
        self.assertEqual(base.output_files, [])
        incoming = xlsx(
            ["号码", "姓名", "国家", "备注"],
            [
                ["+14155552671", "旧号码", "美国", "不导出"],
                ["+33612345678", "重复甲", "法国", "第一行"],
                ["+33 6 12 34 56 78", "重复乙", "法国", "第二行"],
                ["+12025550123", "新用户", "美国", "保留全部字段"],
                ["+4930123456", "德国用户", "德国", "另一国家"],
            ],
        )
        result = self.process(incoming, "new.xlsx", "check", 2)
        self.assertEqual((result.total_rows, result.repeated_rows, result.existing_rows, result.new_rows), (5, 2, 1, 2))
        self.assertEqual({country for country, _ in result.output_files}, {"US", "DE"})
        exported = {country: rows_from(path) for country, path in result.output_files}
        self.assertEqual(exported["US"][1], ("+12025550123", "新用户", "美国", "保留全部字段"))
        self.assertEqual(exported["DE"][1], ("+4930123456", "德国用户", "德国", "另一国家"))
        self.assertEqual(len(exported["US"]), 2)
        self.assertEqual(len(exported["DE"]), 2)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM known_numbers").fetchone()[0], 4)

    def test_country_prefix_and_unknown_are_separate_files(self):
        result = self.process(b"+14155552671\n+442079460958\n9876543210\n", "numbers.txt", "check", 3)
        self.assertEqual(result.new_rows, 3)
        self.assertEqual({country for country, _ in result.output_files}, {"US", "GB", "未识别国家"})
        for _, path in result.output_files:
            self.assertEqual(rows_from(path)[0], ("号码",))

    def test_local_numbers_use_country_and_match_equivalent_international_number(self):
        original = xlsx(["号码", "国家"], [["4155552671", "美国"]])
        self.process(original, "base.xlsx", "base", 20)
        incoming = xlsx(
            ["号码", "国家", "姓名"],
            [["+14155552671", "美国", "已存在"], ["1234567890", "美国", "美国新号"], ["1234567890", "德国", "德国新号"]],
        )
        result = self.process(incoming, "new.xlsx", "check", 21)
        self.assertEqual(result.existing_rows, 1)
        self.assertEqual(result.repeated_rows, 0)
        self.assertEqual(result.new_rows, 2)
        self.assertEqual({country for country, _ in result.output_files}, {"US", "DE"})

    def test_replayed_telegram_update_returns_same_files(self):
        data = b"+14155552671\n"
        first = self.process(data, "numbers.txt", "check", 4)
        second = self.process(data, "numbers.txt", "check", 4)
        self.assertEqual(second.new_rows, 1)
        self.assertEqual(second.output_files, first.output_files)
        self.assertEqual(latest_result(self.db, -100, 7), first.output_files)

    def test_csv_custom_column_and_text_encodings(self):
        data = "客户电话,名字,来源\n+14155552671,甲,广告\n".encode("utf-8-sig")
        result = process_file(self.db, data, "data.csv", "check", self.root / "results", -100, 7, "客户电话", 5)
        self.assertEqual(result.new_rows, 1)
        self.assertEqual(rows_from(result.output_files[0][1])[1][1:], ("甲", "广告"))
        self.assertEqual(normalize_phone("00 44 (20) 7946-0958"), "442079460958")

    def test_csv_keeps_other_fields_including_newlines_and_spaces(self):
        data = '号码,姓名,备注\n+14155552671, 甲 ,"第一行\n第二行"\n'.encode()
        result = self.process(data, "data.csv", "check", 7)
        self.assertEqual(rows_from(result.output_files[0][1])[1], ("+14155552671", " 甲 ", "第一行\n第二行"))

    def test_multiline_customer_template_keeps_fields_and_excludes_all_duplicates(self):
        data = """编号1
客户电话号码: 6591234567
客户姓名: 甲
损失金额: 10000美元 - 20000美元
发生欺诈的平台: Azimuth Management

编号2
客户电话号码: +65 9123 4567
客户姓名: 乙
损失金额: $4111.75
发生欺诈的平台: Carousell

编号3
客户电话号码: +65 8765 4321
客户姓名: 丙
损失金额: 7000美元
发生欺诈的平台:
股票市场投资
其他客户信息: 备注原文
地址: 1 JURONG WEST CENTRAL
SINGAPORE 648886
""".encode()
        result = self.process(data, "records.txt", "check", 30)
        self.assertEqual((result.total_rows, result.repeated_rows, result.new_rows), (3, 2, 1))
        self.assertEqual(result.output_files[0][0], "SG")
        exported = rows_from(result.output_files[0][1])
        self.assertEqual(len(exported), 2)
        headers = exported[0]
        retained = dict(zip(headers, exported[1]))
        self.assertEqual(retained["客户姓名"], "丙")
        self.assertEqual(retained["损失金额"], "7000美元")
        self.assertEqual(retained["发生欺诈的平台"], "股票市场投资")
        self.assertEqual(retained["地址"], "1 JURONG WEST CENTRAL\nSINGAPORE 648886")

    def test_ambiguous_columns_are_rejected(self):
        data = xlsx(["姓名", "账号"], [["甲", "1234567890"]])
        with self.assertRaises(ImportErrorForUser):
            self.process(data, "data.xlsx", "check", 6)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM known_numbers").fetchone()[0], 0)

    def test_bot_only_responds_to_selected_group(self):
        class FakeTelegram:
            def __init__(self):
                self.sent = []

            def download(self, file_id, file_size):
                return b"+14155552671\n"

            def send_message(self, chat_id, text):
                self.sent.append(("message", chat_id, text))

            def send_document(self, chat_id, path, filename, caption):
                self.sent.append(("document", chat_id, filename))

            def call(self, method, payload):
                if method == "getChatMember":
                    return {"status": "member"}
                raise AssertionError(method)

        with patch.dict(os.environ, {
            "NUMBER_BOT_TOKEN": "test-only",
            "NUMBER_BOT_ALLOWED_CHAT_ID": "-100",
            "NUMBER_BOT_DATA_DIR": str(self.root / "bot-data"),
        }):
            bot = Bot()
        self.addCleanup(bot.db.close)
        fake = FakeTelegram()
        bot.telegram = fake
        bot._set_mode(7, "base")
        self.assertEqual(bot._get_mode(7), "base")
        bot._clear_mode(7)
        self.assertEqual(bot._get_mode(7), "check")
        message = {
            "chat": {"id": -101, "type": "supergroup"},
            "from": {"id": 7, "is_bot": False},
            "document": {"file_id": "f", "file_name": "numbers.txt", "file_size": 12},
        }
        with patch("bot.time.sleep"):
            bot.handle({"update_id": 10, "message": message})
            self.assertEqual(fake.sent, [])
            message["chat"]["id"] = -100
            bot.handle({"update_id": 11, "message": message})
            self.assertEqual([item[0] for item in fake.sent], ["message", "document"])
            self.assertEqual(bot.db.execute("SELECT COUNT(*) FROM known_numbers").fetchone()[0], 1)
            bot.handle({"update_id": 11, "message": message})
        self.assertEqual(bot.db.execute("SELECT COUNT(*) FROM known_numbers").fetchone()[0], 1)
        self.assertEqual([item[0] for item in fake.sent], ["message", "document"])


if __name__ == "__main__":
    unittest.main()
